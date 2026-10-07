from __future__ import annotations

import re
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime
from getpass import getpass
from pathlib import Path
from typing import Iterable

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_URL = "https://pogodaiklimat.ru/summary.php"
FIRST_YEAR = 1971
REQUEST_TIMEOUT_SECONDS = 30
REQUEST_DELAY_SECONDS = 0.4
DATE_PATTERN = re.compile(r"^[0-9]{2}\.[0-9]{2}\.[0-9]{4}$")
NUMBER_PATTERN = re.compile(r"[-−+]?\d+(?:[.,]\d+)?")


class ScrapeError(RuntimeError):
    """Raised when the source page does not match the expected structure."""


@dataclass(frozen=True)
class DailyRecord:
    station_number: str
    station_name: str
    observation_date: date
    daily_min_c: float | None
    daily_max_c: float | None
    daily_avg_c: float | None


def normalized_header(value: str) -> str:
    value = value.replace("ё", "е").replace("Ё", "Е").replace("\xa0", " ")
    value = re.sub(r"\s+", " ", value).strip().casefold()
    value = re.sub(r"\s*\([^)]*\)", "", value)
    return re.sub(r"\s+", " ", value).strip()


def row_date(row) -> date | None:
    for cell in row.find_all(["td", "th"], recursive=False):
        text = cell.get_text(" ", strip=True)
        if DATE_PATTERN.fullmatch(text):
            try:
                return datetime.strptime(text, "%d.%m.%Y").date()
            except ValueError:
                return None
    return None


def header_column_indices(header_rows: Iterable) -> dict[str, int]:
    """Expand row/col spans into a logical grid and locate the three leaf labels."""
    grid: list[list[str]] = []
    for row_index, row in enumerate(header_rows):
        if len(grid) <= row_index:
            grid.append([])
        column = 0
        for cell in row.find_all(["td", "th"], recursive=False):
            while column < len(grid[row_index]) and grid[row_index][column]:
                column += 1
            text = cell.get_text(" ", strip=True)
            colspan = max(1, int(cell.get("colspan", 1)))
            rowspan = max(1, int(cell.get("rowspan", 1)))
            for r in range(row_index, row_index + rowspan):
                while len(grid) <= r:
                    grid.append([])
                if len(grid[r]) < column + colspan:
                    grid[r].extend([""] * (column + colspan - len(grid[r])))
                for c in range(column, column + colspan):
                    grid[r][c] = text
            column += colspan

    wanted = {"т ср": "daily_avg", "т мин": "daily_min", "т макс": "daily_max"}
    found: dict[str, int] = {}
    for row in grid:
        for column, text in enumerate(row):
            label = normalized_header(text)
            if label in wanted:
                found[wanted[label]] = column
    missing = set(wanted.values()) - set(found)
    if missing:
        raise ScrapeError(
            "Could not identify temperature columns in the page header "
            f"(missing: {', '.join(sorted(missing))}). The website may have changed."
        )
    return found


def parse_temperature(cell) -> float | None:
    text = cell.get_text(" ", strip=True)
    if not text:
        return None
    match = NUMBER_PATTERN.search(text)
    if not match:
        return None
    return float(match.group(0).replace("−", "-").replace(",", "."))


def deduplicate_records(records: list[DailyRecord]) -> tuple[list[DailyRecord], list[date]]:
    """Drop exact daily duplicates and return dates with conflicting temperatures."""
    unique_records: list[DailyRecord] = []
    seen_records: set[tuple[date, float | None, float | None, float | None]] = set()
    temperatures_by_date: dict[date, set[tuple[float | None, float | None, float | None]]] = {}

    for record in records:
        temperatures = (record.daily_min_c, record.daily_max_c, record.daily_avg_c)
        temperatures_by_date.setdefault(record.observation_date, set()).add(temperatures)
        signature = (record.observation_date, *temperatures)
        if signature in seen_records:
            continue
        seen_records.add(signature)
        unique_records.append(record)

    suspect_dates = sorted(
        day for day, temperature_values in temperatures_by_date.items() if len(temperature_values) > 1
    )
    return unique_records, suspect_dates


def parse_year_page(html: str, station_number: str, year: int) -> list[DailyRecord]:
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table", id="summary_table")
    if table is None:
        page_text = soup.get_text(" ", strip=True).casefold()
        if "доступен только для зарегистрированных пользователей" in page_text:
            raise ScrapeError(
                "Pogodaiklimat did not grant access to the summary table. Check that the username and password are correct and that the account is registered."
            )
        raise ScrapeError(
            f"Year {year}: the page did not contain the weather table. "
            "The site may require sign-in or may have changed its page."
        )

    rows = table.find_all("tr")
    first_data_index = next((i for i, row in enumerate(rows) if row_date(row)), None)
    if first_data_index is None:
        header_rows = rows
    else:
        header_rows = rows[:first_data_index]
    columns = header_column_indices(header_rows)

    records: list[DailyRecord] = []
    for row in rows:
        observation_date = row_date(row)
        if observation_date is None or observation_date.year != year:
            continue
        cells = row.find_all(["td", "th"], recursive=False)
        if len(cells) <= max(columns.values()) or len(cells) < 3:
            continue

        row_station = cells[0].get_text(" ", strip=True)
        if row_station != station_number:
            continue
        station_name = cells[1].get_text(" ", strip=True)
        # A date row with no values in any of the requested fields has no useful record.
        temps = {
            key: parse_temperature(cells[index])
            for key, index in columns.items()
        }
        if all(value is None for value in temps.values()):
            continue
        records.append(
            DailyRecord(
                station_number=station_number,
                station_name=station_name,
                observation_date=observation_date,
                daily_min_c=temps["daily_min"],
                daily_max_c=temps["daily_max"],
                daily_avg_c=temps["daily_avg"],
            )
        )
    return records


def fetch_year(
    session: requests.Session,
    station_number: str,
    year: int,
    month: int | None = None,
) -> list[DailyRecord]:
    params = {"y": year, "id": station_number}
    if month is not None:
        params["m"] = month
    try:
        response = session.get(
            BASE_URL,
            params=params,
            timeout=REQUEST_TIMEOUT_SECONDS,
        )
    except requests.RequestException as exc:
        period = f"{year}-{month:02d}" if month is not None else str(year)
        raise ScrapeError(f"Period {period}: request failed: {exc}") from exc

    if response.status_code == 404:
        print(f"  Year {year}: page not found (404); skipped.")
        return []
    try:
        response.raise_for_status()
    except requests.HTTPError as exc:
        raise ScrapeError(f"Year {year}: HTTP error: {exc}") from exc
    response.encoding = response.apparent_encoding or response.encoding
    records = parse_year_page(response.text, station_number, year)
    if month is not None:
        records = [record for record in records if record.observation_date.month == month]
    return records


def authenticate(session: requests.Session, username: str, password: str) -> None:
    """Sign in to the site; the caller owns the session and credentials are not persisted."""
    response = session.post(
        "https://pogodaiklimat.ru/login.php",
        data={"submit-login": "post", "username": username, "password": password},
        timeout=REQUEST_TIMEOUT_SECONDS,
    )
    response.raise_for_status()


def write_workbook(records: list[DailyRecord], station_number: str, output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Daily temperatures"
    headers = [
        "Station number",
        "Station name",
        "Year",
        "Month",
        "Day",
        "min",
        "max",
        "avg",
    ]
    sheet.append(headers)
    for record in records:
        sheet.append(
            [
                record.station_number,
                record.station_name,
                record.observation_date.year,
                record.observation_date.month,
                record.observation_date.day,
                record.daily_min_c,
                record.daily_max_c,
                record.daily_avg_c,
            ]
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for column, width in {
        "A": 16,
        "B": 24,
        "C": 10,
        "D": 10,
        "E": 10,
        "F": 16,
        "G": 16,
        "H": 30,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    if records:
        table = Table(displayName="DailyWeather", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def prompt_station_number() -> str:
    while True:
        station_number = input("Enter the 5-digit weather station number: ").strip()
        if re.fullmatch(r"[0-9]{5}", station_number):
            return station_number
        print("Please enter exactly 5 digits (for example: 24688).")


def main() -> int:
    station_number = prompt_station_number()
    username = input("Pogodaiklimat username: ").strip()
    password = getpass("Pogodaiklimat password: ")
    current_year = date.today().year
    output_path = Path.cwd() / f"pogodaiklimat_{station_number}_daily_weather.xlsx"
    print(f"Fetching station {station_number}, {FIRST_YEAR}–{current_year}…")

    all_records: list[DailyRecord] = []
    try:
        with requests.Session() as session:
            session.headers.update(
                {"User-Agent": "PogodaiklimatDailyWeatherScraper/1.0 (personal data export)"}
            )
            authenticate(session, username, password)
            password = ""
            for year in range(FIRST_YEAR, current_year + 1):
                print(f"[{year}] Loading…", end=" ", flush=True)
                year_records = fetch_year(session, station_number, year)
                all_records.extend(year_records)
                print(f"{len(year_records)} daily records")
                if year < current_year:
                    time.sleep(REQUEST_DELAY_SECONDS)
    except (ScrapeError, requests.RequestException) as exc:
        print(f"\nStopped: {exc}", file=sys.stderr)
        print("No workbook was written, so it cannot be mistaken for a complete export.", file=sys.stderr)
        return 1

    all_records, suspect_dates = deduplicate_records(all_records)
    all_records.sort(key=lambda record: record.observation_date)
    if suspect_dates:
        print("Suspect dates with conflicting temperatures: " + ", ".join(day.strftime("%d.%m.%Y") for day in suspect_dates))
    write_workbook(all_records, station_number, output_path)
    print(f"Done: {len(all_records)} records saved to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

